"""Load the Momentuum Blue workbook into SQLite.

HR-Input, Input - Trainers, and Master-Input are the typed sheets.
Dashboard tabs are formulas over those three, so they are not imported.
"""

import logging
from datetime import date, datetime

from app.database import EXCEL_PATH, SessionLocal
from app.models import Person, PipelineWeek, TrainingWeek

logger = logging.getLogger("uvicorn.error")

HR_FIELDS = (
    "applications_total",
    "screening_pending",
    "screening_cleared",
    "l1_pending",
    "l1_cleared",
    "berribot_applied",
    "berribot_cleared",
    "l2_pending",
    "l2_cleared",
    "l3_pending",
    "l3_cleared",
    "cto_pending",
    "cto_cleared",
    "offer_extended",
    "offer_accepted",
    "joined",
    "on_hold",
    "rejected",
    "slots_needed",
)


def _date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _int(value):
    if value is None or value == "":
        return None
    return int(value)


def _float(value):
    if value is None or value == "":
        return None
    return float(value)


def _text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def load(db, path=EXCEL_PATH):
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)

    hr = workbook["HR-Input"]
    for row in hr.iter_rows(min_row=2, max_col=24, values_only=True):
        week_ending = _date(row[0])
        cohort = _text(row[1])
        role = _text(row[2])
        if week_ending is None or cohort is None or role is None:
            continue
        counts = [_int(value) for value in row[4:23]]
        payload = dict(zip(HR_FIELDS, counts))
        db.add(
            PipelineWeek(
                week_ending=week_ending,
                cohort=cohort,
                role=role,
                source=_text(row[3]),
                notes=_text(row[23]),
                **payload,
            )
        )

    trainers = workbook["Input - Trainers"]
    for row in trainers.iter_rows(min_row=2, max_col=6, values_only=True):
        week = _int(row[0])
        if week is None:
            continue
        db.add(
            TrainingWeek(
                week=week,
                hands_on=_float(row[1]),
                proctored=_float(row[2]),
                overall=_float(row[3]),
                active_students=_int(row[4]),
                topic=_text(row[5]),
            )
        )

    roster = workbook["Master-Input"]
    for row in roster.iter_rows(min_row=2, max_col=11, values_only=True):
        name = _text(row[0])
        if name is None:
            continue
        db.add(
            Person(
                name=name,
                location=_text(row[1]) or "Unspecified",
                career_stage=_text(row[2]) or "Unspecified",
                start_date=_date(row[3]),
                source=_text(row[4]),
                client=_text(row[5]),
                client_start=_date(row[6]),
                client_end=_date(row[7]),
                hourly_rate=_float(row[8]),
                hours_per_week=_float(row[9]),
            )
        )

    workbook.close()


def seed_if_empty():
    db = SessionLocal()
    try:
        if db.query(Person).count():
            return
        if not EXCEL_PATH.exists():
            logger.warning("Workbook not found at %s; starting with an empty database", EXCEL_PATH)
            return
        load(db)
        db.commit()
        logger.info(
            "Seeded %s people, %s hiring rows, %s academy weeks",
            db.query(Person).count(),
            db.query(PipelineWeek).count(),
            db.query(TrainingWeek).count(),
        )
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
